"""Data dumps: Excel workbook, CSV zip, JSON full dump (+ restore)."""
import csv
import enum
import io
import json
import zipfile
from datetime import date, datetime

from typing import Any

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.database import Base
from app.dates import display_dates_in_text, format_date, format_datetime, format_ist_stamp
from app.models import entities  # noqa: F401  (registers every table on Base.metadata)

DUMP_FORMAT = "nest5-ledger-dump"
DUMP_FORMAT_VERSION = 1

# Secrets stay out of spreadsheet/CSV exports; the JSON dump keeps them so a restore
# brings logins back exactly as they were.
REDACTED_COLUMNS = {("users", "password_hash"), ("password_reset_tokens", "token_hash")}

_HEADER_WORDS = {"id": "ID", "json": "JSON", "vpa": "UPI ID", "upi": "UPI"}


def tables() -> list[sa.Table]:
    return list(Base.metadata.sorted_tables)


def schema_version(db: Session) -> str | None:
    try:
        return db.execute(sa.text("SELECT version_num FROM alembic_version")).scalar()
    except Exception:  # noqa: BLE001 — table absent on DBs created without alembic
        db.rollback()
        return None


def fetch_rows(db: Session, table: sa.Table, *, reports: bool = False) -> list[dict[str, Any]]:
    """`reports=True` (Excel / CSV) leaves out cancelled / rejected / reverted transactions; the
    JSON backup keeps every row (expenses carry `status` + `deleted_at`, rejected owner payments
    live in `submission_rejections`)."""
    order = list(table.primary_key.columns) or list(table.columns)
    q = sa.select(table).order_by(*order)
    if reports:
        if table.name == "expenses" and "status" in table.c:
            q = q.where(sa.or_(table.c.status.is_(None), table.c.status != "rejected"), table.c.deleted_at.is_(None))
        elif table.name == "settlements" and "reverted_at" in table.c:
            counted = sa.select(sa.text("id")).select_from(sa.table("expenses")).where(
                sa.or_(sa.column("status").is_(None), sa.column("status") != "rejected"),
                sa.column("deleted_at").is_(None),
            )
            q = q.where(table.c.reverted_at.is_(None), table.c.expense_id.in_(counted))
        elif table.name == "submission_rejections":
            q = q.where(sa.false())
        elif table.name == "users" and "deleted_at" in table.c:
            q = q.where(table.c.deleted_at.is_(None))
    return [dict(r._mapping) for r in db.execute(q)]


def row_counts(db: Session) -> dict[str, int]:
    return {
        t.name: int(db.execute(sa.select(sa.func.count()).select_from(t)).scalar() or 0)
        for t in tables()
    }


def readable_header(column: str) -> str:
    words = column.split("_")
    if words[0] == "is" and len(words) > 1:
        words = words[1:] + ["?"]
    out = [_HEADER_WORDS.get(w, w) for w in words]
    text = " ".join(out).replace(" ?", "?")
    return text[:1].upper() + text[1:]


def readable_table(name: str) -> str:
    text = name.replace("_", " ")
    return (text[:1].upper() + text[1:])[:31]


def _visible_columns(table: sa.Table) -> list[sa.Column]:
    return [c for c in table.columns if (table.name, c.name) not in REDACTED_COLUMNS]


def _raw(value: Any) -> Any:
    """Value as SQLite stores it (enums by name, dates ISO)."""
    if isinstance(value, enum.Enum):
        return value.name
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


# ---- Excel ------------------------------------------------------------------------


def build_xlsx(db: Session) -> bytes:
    wb = Workbook()
    info = wb.active
    info.title = "About"
    info.append(["Nest5 Ledger export"])
    info["A1"].font = Font(bold=True, size=14)
    info.append(["Exported at", format_ist_stamp()])
    info.append(["Schema version", schema_version(db) or "—"])
    info.append([])
    info.append(["Sheet", "Rows"])
    info["A5"].font = info["B5"].font = Font(bold=True)
    for t in tables():
        cols = _visible_columns(t)
        ws = wb.create_sheet(readable_table(t.name))
        ws.append([readable_header(c.name) for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        rows = fetch_rows(db, t, reports=True)
        for row in rows:
            out = []
            for c in cols:
                v = row[c.name]
                if isinstance(v, enum.Enum):
                    v = v.value
                elif isinstance(v, bool):
                    v = "Yes" if v else "No"
                elif c.name.endswith("_json") and isinstance(v, str):
                    v = display_dates_in_text(v)
                out.append(v)
            ws.append(out)
        for idx, c in enumerate(cols, start=1):
            letter = get_column_letter(idx)
            if isinstance(c.type, sa.DateTime):
                fmt = "dd-mm-yyyy hh:mm"
            elif isinstance(c.type, sa.Date):
                fmt = "dd-mm-yyyy"
            else:
                fmt = None
            if fmt:
                for cell in ws[letter][1:]:
                    cell.number_format = fmt
            width = max([len(readable_header(c.name))] + [len(str(r[c.name] or "")) for r in rows[:200]])
            ws.column_dimensions[letter].width = min(max(width + 2, 10), 60)
        info.append([ws.title, len(rows)])
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- CSV zip ----------------------------------------------------------------------


def build_csv_zip(db: Session) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for t in tables():
            cols = _visible_columns(t)
            text = io.StringIO()
            writer = csv.writer(text)
            writer.writerow([c.name for c in cols])
            for row in fetch_rows(db, t, reports=True):
                out = []
                for c in cols:
                    v = row[c.name]
                    if isinstance(v, datetime):
                        v = format_datetime(v)
                    elif isinstance(v, date):
                        v = format_date(v)
                    else:
                        v = _raw(v)
                    if c.name.endswith("_json") and isinstance(v, str):
                        v = display_dates_in_text(v)
                    if isinstance(v, bool):
                        v = int(v)
                    out.append("" if v is None else v)
                writer.writerow(out)
            zf.writestr(f"{t.name}.csv", text.getvalue())
        zf.writestr(
            "README.txt",
            "Nest5 Ledger CSV export — one file per table, header row = column names.\n"
            "Dates are dd-mm-yyyy (timestamps dd-mm-yyyy HH:MM, as stored: UTC), booleans 1/0,\n"
            "empty cell = NULL. Password hashes are omitted; use the JSON dump for a full restore.\n",
        )
    return buf.getvalue()


# ---- JSON dump / restore ----------------------------------------------------------


def build_json_dump(db: Session) -> dict:
    data: dict[str, Any] = {
        "format": DUMP_FORMAT,
        "format_version": DUMP_FORMAT_VERSION,
        "schema_version": schema_version(db),
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "tables": {},
    }
    for t in tables():
        rows = fetch_rows(db, t)
        data["tables"][t.name] = {
            "columns": [c.name for c in t.columns],
            "rows": [{k: _raw(v) for k, v in r.items()} for r in rows],
        }
    data["row_counts"] = {name: len(v["rows"]) for name, v in data["tables"].items()}
    return data


def dump_json_bytes(db: Session) -> bytes:
    return json.dumps(build_json_dump(db), ensure_ascii=False, indent=1).encode("utf-8")


class RestoreError(ValueError):
    pass


def validate_dump(data: Any) -> dict[str, int]:
    """Check a parsed dump and return its per-table row counts."""
    if not isinstance(data, dict) or data.get("format") != DUMP_FORMAT:
        raise RestoreError("Not a Nest5 ledger JSON dump.")
    if int(data.get("format_version", 0)) > DUMP_FORMAT_VERSION:
        raise RestoreError("This dump was made by a newer version of the app.")
    known = {t.name: t for t in tables()}
    dump_tables = data.get("tables")
    if not isinstance(dump_tables, dict):
        raise RestoreError("Dump has no tables.")
    counts = {}
    for name, payload in dump_tables.items():
        if name not in known:
            raise RestoreError(f"Unknown table in dump: {name}")
        rows = payload.get("rows", []) if isinstance(payload, dict) else None
        if not isinstance(rows, list):
            raise RestoreError(f"Table {name} has no row list.")
        cols = set(known[name].columns.keys())
        for row in rows:
            extra = set(row) - cols
            if extra:
                raise RestoreError(f"Table {name} has columns this app does not know: {sorted(extra)}")
        counts[name] = len(rows)
    return counts


def _from_raw(column: sa.Column, value: Any) -> Any:
    if value is None:
        return None
    ctype = column.type
    if isinstance(ctype, sa.DateTime):
        return datetime.fromisoformat(value)
    if isinstance(ctype, sa.Date):
        return date.fromisoformat(value[:10])
    if isinstance(ctype, sa.Enum) and ctype.enum_class is not None:
        try:
            return ctype.enum_class[value]
        except KeyError:
            return ctype.enum_class(value)
    if isinstance(ctype, sa.Boolean):
        return bool(value)
    return value


def restore_dump(db: Session, data: dict) -> dict[str, int]:
    """Replace table contents with the dump. Tables missing from the dump are left untouched.

    Runs on the caller's session; the caller commits (or rolls back on error).
    """
    validate_dump(data)
    present = [t for t in tables() if t.name in data["tables"]]
    for t in reversed(present):
        db.execute(t.delete())
    for t in present:
        rows = data["tables"][t.name].get("rows", [])
        if rows:
            converted = [{k: _from_raw(t.c[k], v) for k, v in row.items()} for row in rows]
            db.execute(t.insert(), converted)
    db.flush()
    return {t.name: len(data["tables"][t.name].get("rows", [])) for t in present}
