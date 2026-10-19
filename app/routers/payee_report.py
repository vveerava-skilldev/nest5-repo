"""Admin-only Payee / UPI audit report: read-only, filterable, exportable (CSV / Excel / PDF)."""
import csv
import io
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import HTMLResponse, Response
from sqlalchemy.orm import Session

from app.database import get_db
from app.dates import format_date, format_ist, format_ist_stamp, ist_now
from app.routers.web import templates
from app.services import payee_report as report_svc
from app.services.access import is_admin

router = APIRouter()

REPORT_URL = "/admin/payee-report"


def _admin(request: Request) -> None:
    if not is_admin(request):
        raise HTTPException(403, detail="This report is for administrators only.")


def _filter_text(f: report_svc.Filters) -> str:
    bits = []
    for label, val in (("UPI ID", f.vpa_text or f.vpa), ("mobile", f.mobile), ("payee", f.payee),
                       ("settings version", f.version), ("from", f.date_from and format_date(f.date_from)),
                       ("to", f.date_to and format_date(f.date_to)),
                       ("type", f.type if f.type != "all" else ""), ("mode", f.mode if f.mode != "all" else ""),
                       ("status", report_svc.STATUSES[f.status]), ("flat", f.flat)):
        if val:
            bits.append(f"{label}: {val}")
    if f.include:
        bits.append("rejected and reverted rows included in totals")
    return "; ".join(bits)


@router.get(REPORT_URL, response_class=HTMLResponse)
def payee_report_page(request: Request, db: Session = Depends(get_db)):
    _admin(request)
    f = report_svc.parse(request.query_params)
    data = report_svc.build(db, f)
    return templates.TemplateResponse(
        request, "admin_payee_report.html",
        {"request": request, "f": f, "d": data, "statuses": report_svc.STATUSES, "types": report_svc.TYPES,
         "modes": report_svc.MODES, "no_upi": report_svc.NO_UPI, "export_query": _query(f), "page_size": report_svc.PAGE_SIZE,
         "filter_text": _filter_text(f)},
    )


def _query(f: report_svc.Filters) -> str:
    from urllib.parse import urlencode

    return urlencode(f.as_query())


def _filtered(request: Request, db: Session):
    _admin(request)
    f = report_svc.parse(request.query_params)
    data = report_svc.build(db, f)
    return f, data


def _download(content: bytes, name: str, mime: str) -> Response:
    return Response(content, media_type=mime, headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _stamp() -> str:
    return ist_now().strftime("%Y%m%d-%H%M")


@router.get(REPORT_URL + "/export.csv")
def payee_report_csv(request: Request, db: Session = Depends(get_db)):
    f, data = _filtered(request, db)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(report_svc.COLUMNS)
    w.writerows(report_svc.export_rows(data["rows"], f.include))
    return _download(buf.getvalue().encode("utf-8"), f"payee-report-{_stamp()}.csv", "text/csv")


def _summary_tables(data: dict) -> tuple[list[list], list[list]]:
    by_payee = [["Paid to", "Payments counted", "Total received (Rs.)", "Excluded rows", "Excluded amount (Rs.)"]]
    for g in data["summary"]["by_payee"]:
        by_payee.append([g["payee"], g["count"], g["total"], g["excluded"], g["excluded_total"]])
    by_payee.append(["All", data["summary"]["count"], data["summary"]["total"], "", ""])
    periods = [["Version", "UPI ID", "Mobile", "Payee", "From (IST)", "To (IST)", "UPI count", "UPI received (Rs.)",
                "Cash/bank count", "Cash/bank (Rs.)", "Account balance when set (Rs.)", "Total balance when set (Rs.)",
                "Account balance at next change (Rs.)", "Bank balance typed"]]
    for p in data["summary"]["periods"]:
        periods.append([
            p["id"], p["vpa"], p["mobile"], p["payee"], format_ist_stamp(p["from"]), format_ist_stamp(p["to"]) if p["to"] else "now",
            p["upi_count"], p["upi_total"], p["other_count"], p["other_total"],
            "" if p["fund_at_start"] is None else p["fund_at_start"], "" if p["total_at_start"] is None else p["total_at_start"],
            "" if p["fund_at_end"] is None else p["fund_at_end"], p["bank_entered"] or "",
        ])
    return by_payee, periods


@router.get(REPORT_URL + "/export.xlsx")
def payee_report_xlsx(request: Request, db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    f, data = _filtered(request, db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Payments"
    ws.append(["Payee / UPI audit report", f"Generated {format_ist_stamp()}"])
    ws.append(["Filters", _filter_text(f)])
    ws.append([])
    ws.append(list(report_svc.COLUMNS))
    for c in ws[4]:
        c.font = Font(bold=True)
    for row in report_svc.export_rows(data["rows"], f.include):
        ws.append(row)
    by_payee, periods = _summary_tables(data)
    for title, table in (("By payee", by_payee), ("By settings version", periods)):
        sheet = wb.create_sheet(title)
        for row in table:
            sheet.append(row)
        for c in sheet[1]:
            c.font = Font(bold=True)
    out = io.BytesIO()
    wb.save(out)
    return _download(out.getvalue(), f"payee-report-{_stamp()}.xlsx",
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.get(REPORT_URL + "/export.pdf")
def payee_report_pdf(request: Request, db: Session = Depends(get_db)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    f, data = _filtered(request, db)
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=6.5, leading=8)

    def tbl(rows, widths=None):
        t = Table([[Paragraph(escape(str(c)), cell) for c in r] for r in rows], colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    by_payee, periods = _summary_tables(data)
    story = [Paragraph("Payee / UPI audit report", styles["Title"]),
             Paragraph(escape(f"Generated {format_ist_stamp()}. Filters: {_filter_text(f) or 'none'}"), styles["BodyText"]),
             Paragraph("Totals count paid rows only (rejected and reverted rows are listed but excluded unless included).", styles["BodyText"]),
             Spacer(1, 4 * mm), Paragraph("By payee", styles["Heading3"]), tbl(by_payee), Spacer(1, 4 * mm),
             Paragraph("By settings version (with the balance the app recorded at each change)", styles["Heading3"]),
             tbl([r[:2] + r[4:12] for r in periods]), Spacer(1, 4 * mm), Paragraph("Payments", styles["Heading3"]),
             tbl([list(report_svc.COLUMNS)] + report_svc.export_rows(data["rows"], f.include))]
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=8 * mm, rightMargin=8 * mm, topMargin=8 * mm,
                      bottomMargin=8 * mm).build(story)
    return _download(out.getvalue(), f"payee-report-{_stamp()}.pdf", "application/pdf")
