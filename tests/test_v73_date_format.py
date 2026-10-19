"""v7.3 (in place): every date the app shows is dd-mm-yyyy; ISO stays only in form values, URLs, DB and the JSON backup."""
import json
import shutil
import subprocess
from datetime import date, datetime, timezone
from pathlib import Path

import pytest

from app import dates
from tests.test_v71_empty_and_review import run_scratch

ROOT = Path(__file__).resolve().parents[1]


def test_central_helpers():
    assert dates.format_date(date(2026, 10, 5)) == "05-10-2026"
    assert dates.format_date("2026-01-02") == "02-01-2026" and dates.format_date(None) == "—"
    assert dates.format_date(datetime(2026, 10, 5, 23, 59)) == "05-10-2026"
    assert dates.format_datetime(datetime(2026, 10, 5, 9, 7, 33)) == "05-10-2026 09:07"
    assert dates.format_weekday_date(date(2026, 10, 5)) == "Mon 05-10-2026"
    assert dates.due_hint(date(2026, 10, 5)) == "due 05-10-2026" and dates.due_hint(None) == ""
    assert dates.format_date_or_blank(None) == "" and dates.format_date_or_blank(date(2026, 3, 1)) == "01-03-2026"
    assert dates.format_ist(datetime(2026, 10, 19, 11, 7, tzinfo=timezone.utc)) == "Mon 19-10-2026, 4:37 PM IST"
    assert dates.format_ist_stamp(datetime(2026, 10, 19, 11, 7, tzinfo=timezone.utc)) == "19-10-2026 16:37 IST"
    assert dates.display_dates_in_text('{"paid_on": "2026-10-05", "at": "2026-10-05T09:07:33.123456", "n": 12}') == \
        '{"paid_on": "05-10-2026", "at": "05-10-2026 09:07", "n": 12}'
    assert dates.display_dates_in_text(None) == "" and dates.display_dates_in_text("no dates 2026-10") == "no dates 2026-10"
    # form values / DB helpers stay ISO
    assert dates.date_to_iso(date(2026, 10, 5)) == "2026-10-05" and dates.parse_iso_date("2026-10-05") == date(2026, 10, 5)


@pytest.mark.skipif(shutil.which("node") is None, reason="node not installed")
def test_browser_clock_script_matches_the_server_format(tmp_path):
    script = tmp_path / "clock.js"
    script.write_text(
        "const fixed = Date.UTC(2026, 9, 19, 11, 7, 0); const RealDate = Date;\n"
        "global.Date = class extends RealDate { constructor(...a) { super(...(a.length ? a : [fixed])); } static now() { return fixed; } };\n"
        "const el = { textContent: '' }; global.document = { getElementById: () => el };\n"
        "global.setTimeout = () => 0;\n"
        f"require({json.dumps(str(ROOT / 'app/static/js/ist-clock.js'))});\nconsole.log(el.textContent);\n")
    ran = subprocess.run(["node", str(script)], capture_output=True, text=True, timeout=30)
    assert ran.stdout.strip() == "Mon 19-10-2026, 4:37 PM IST" == dates.format_ist(datetime(2026, 10, 19, 11, 7, tzinfo=timezone.utc)), ran.stderr


SCAN = r'''
import csv, io, re, zipfile
import openpyxl
from datetime import date, timedelta
OWNER = "asrk.sastry"
owner, mani = login(OWNER), login("manikyamba")
tflat = db.scalar(select(Flat).where(Flat.number == "101"))
ten = User(username="ten1", display_name="Tenant", role="tenant", is_active=True, must_change_password=False, password_hash="x")
ten.flats = [tflat]; db.add(ten); db.commit(); u.set_password(ten, "pw-pw-pw-1", must_change=False); db.commit()
tenant = login("ten1")
sast = db.scalar(select(User).where(User.username == OWNER)); oflat = sast.flats[0].id
by_no = {f.number: f.id for f in db.scalars(select(Flat)).all()}
octo = db.scalar(select(LedgerPeriod).where(LedgerPeriod.year == 2026, LedgerPeriod.month == 10))
mp = db.scalar(select(MaintenancePayment).where(MaintenancePayment.period_id == octo.id, MaintenancePayment.flat_id == oflat))
admin.post(f"/maintenance/{mp.id}/revert")
today = date.today().isoformat()
admin.post("/projects", data={"name": "Roof", "total_cost": "25000", "start_date": "2026-09-01", "due_date": "2026-09-20"})
admin.post("/corpus/campaigns", data={"name": "Lift", "target_amount": "50000", "due_date": "2026-12-15"})
admin.post("/calendar/events", data={"title": "Owners meeting", "event_type": "Meeting", "event_date": "2026-10-25", "end_date": "2026-10-26", "visible_to_owners": "1"})
admin.post("/admin/recurring/tasks", data={"title": "Lift service", "day_of_month": "22", "due_mode": "by", "interval_months": "1", "notes": ""})
admin.post("/admin/payment-settings/review", data={"upi_vpa": "new.soc@oksbi", "upi_mobile": "9876543210", "payee_name": "New Payee", "reason": "moved", "bank_balance": ""})
admin.post("/admin/payment-settings/confirm", data={"upi_vpa": "new.soc@oksbi", "upi_mobile": "9876543210", "payee_name": "New Payee", "reason": "moved", "bank_balance": "", "confirm": "yes"})
owner.post("/owner/maintenance", data={"flat_id": oflat, "period_id": octo.id, "paid_on": "2026-10-05", "amount": "", "payment_mode": "UPI", "payment_note": "UTR123456789012"})
db.expire_all()
proj = db.scalar(select(Project).where(Project.name == "Roof"))
share = db.scalar(select(ProjectPayment).where(ProjectPayment.project_id == proj.id, ProjectPayment.flat_id == oflat))
owner.post("/owner/projects/pay", data={"share_id": share.id, "amount": "100", "paid_on": "2026-10-06", "payment_mode": "Cash"})
owner.post("/owner/projects/pay", data={"share_id": share.id, "amount": "55", "paid_on": "2026-10-06", "payment_mode": "UPI", "payment_note": "UTR998877665544"})
db.expire_all()
rej = db.scalars(select(ProjectContribution).where(ProjectContribution.amount == 55)).first()
admin.post(f"/admin/submissions/project_contribution/{rej.id}/reject", data={"reason": "not received", "back": "/admin/actions"})
owner.post("/owner/expenses", data={"flat_id": oflat, "description": "Pump repair", "category": "Common", "amount": "1850", "expense_date": "2026-10-04"})
admin.post("/expenses", data={"description": "Bulbs", "category": "Common", "amount": "300", "expense_date": "2026-10-03"})

ISO = re.compile(r"\b\d{4}-\d{2}-\d{2}\b")
MON = "(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*"
OTHER = [re.compile(rf"\b\d{{1,2}} {MON},? \d{{4}}"), re.compile(rf"\b{MON} \d{{1,2}},? \d{{4}}"), re.compile(r"\b\d{1,2}/\d{1,2}/\d{4}\b")]
def visible(html):
    h = re.sub(r"<(script|style)\b.*?</\1>", " ", html, flags=re.S)
    attrs = re.findall(r'\b(?:title|placeholder|alt|aria-label)="([^"]*)"', h)
    return re.sub(r"<[^>]+>", " ", h) + "\n" + "\n".join(attrs)
def leaks(text):
    out = []
    for rx in [ISO] + OTHER:
        out += [text[max(0, m.start() - 40):m.end() + 15].replace("\n", " ") for m in rx.finditer(text)]
    return out
SKIP = re.compile(r"/(logout|revert|settle)\b|/delete|/export|/docs|/redoc|/openapi|/upi/qr|calendar\?m=|page=|statement\.pdf")
def crawl(c, seeds, limit=250):
    seen, todo, bad = set(), list(seeds), {}
    while todo and len(seen) < limit:
        url = todo.pop(0)
        if url in seen or SKIP.search(url):
            continue
        seen.add(url)
        r = c.get(url, follow_redirects=True)
        if r.status_code != 200 or "html" not in r.headers.get("content-type", ""):
            continue
        if leaks(visible(r.text)):
            bad[url] = leaks(visible(r.text))[:3]
        todo += [h.replace("&amp;", "&") for h in re.findall(r'href="(/[^"#]*)"', r.text) if not h.startswith("/static")]
    return bad, len(seen)
ADMIN_SEEDS = ["/", "/?view=simple", "/?period_id=1", "/?period_id=2", "/payments/maintenance", "/payments/corpus", "/payments/projects", "/payments/expenses", "/payments/building",
               "/calendar", "/calendar?m=2026-10", "/schedule", "/audit", "/admin/activity", "/admin/actions", "/admin/payment-settings", "/admin/payment-settings/history",
               "/admin/payee-report?status=all", "/admin/payee-report?date_from=2026-10-01&date_to=2026-10-31", "/admin/recurring", "/settlements", "/corpus", "/projects", "/flats", "/expenses",
               "/admin/users", "/admin/export", "/complaints", "/periods/rollover", "/admin/cancel-requests", "/admin/maintenance-fee", "/calendar/events/new", "/admin/expense-categories", "/account/password"]
OWNER_SEEDS = ["/", "/?view=simple", "/?view=detailed", "/?view=simple&period_id=2&pay=m-2-2", "/payments/maintenance", "/payments/corpus", "/payments/projects", "/payments/expenses", "/payments/building",
               "/calendar", "/schedule", "/owner/submissions", "/corpus", "/projects", "/expenses", "/complaints", "/account/password"]
TENANT_SEEDS = ["/", "/payments/maintenance", "/owner/submissions", "/account/password"]
'''


def test_no_iso_date_leaks_on_main_admin_owner_and_tenant_pages(tmp_path):
    code = r'''
out["admin"] = crawl(admin, ADMIN_SEEDS)
out["owner"] = crawl(owner, OWNER_SEEDS)
out["mani"] = crawl(mani, OWNER_SEEDS)
out["tenant"] = crawl(tenant, TENANT_SEEDS)
# the audit page's raw before / after data is shown dd-mm-yyyy too, while the stored JSON is untouched
page = admin.get("/audit").text
out["audit_raw"] = [bool(re.search(r"<code>[^<]*\b\d\d-\d\d-\d{4}\b", page)), bool(ISO.search(re.sub(r"<[^>]+>", " ", page)))]
db.expire_all()
out["audit_db_iso"] = any(ISO.search(a.after_json or "") for a in db.scalars(select(AuditLog)).all())
out["clock"] = bool(re.search(r'id="ist-clock"[^>]*>(Mon|Tue|Wed|Thu|Fri|Sat|Sun) \d\d-\d\d-\d{4}, \d{1,2}:\d\d (AM|PM) IST<', admin.get("/").text))
out["pages_seen"] = [out["admin"][1], out["owner"][1], out["tenant"][1]]
'''
    out = run_scratch(tmp_path, SCAN + code)
    for who in ("admin", "owner", "mani", "tenant"):
        assert out[who][0] == {}, (who, out[who][0])
    assert out["pages_seen"][0] >= 40 and out["pages_seen"][1] >= 15 and out["pages_seen"][2] >= 3
    assert out["audit_raw"] == [True, False] and out["audit_db_iso"] is True
    assert out["clock"] is True


def test_exports_show_dd_mm_yyyy_but_the_json_backup_stays_iso(tmp_path):
    code = r'''
def sheet_text(wb):
    cells = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    cells.append(c.value)
                elif isinstance(c.value, (date, datetime)):
                    cells.append(("DATE", c.number_format))
    return cells
from datetime import datetime
xl = openpyxl.load_workbook(io.BytesIO(admin.get("/admin/export/excel").content))
cells = sheet_text(xl)
out["xlsx_text_leaks"] = [t for t in cells if isinstance(t, str) and leaks(t)][:3]
out["xlsx_date_formats"] = sorted({t[1] for t in cells if isinstance(t, tuple)})
out["xlsx_about"] = [r[1].value for r in xl["About"].iter_rows(min_row=2, max_row=2)]
z = zipfile.ZipFile(io.BytesIO(admin.get("/admin/export/csv").content))
csv_bad = {n: leaks(z.read(n).decode()) [:2] for n in z.namelist() if n.endswith((".csv", ".txt")) and leaks(z.read(n).decode())}
out["csv_bad"] = csv_bad
mpay = z.read("maintenance_payments.csv").decode().splitlines()
hdr = mpay[0].split(",")
out["csv_paid_on"] = sorted({r.split(",")[hdr.index("paid_on")] for r in mpay[1:] if r.split(",")[hdr.index("paid_on")]})[:3]
out["csv_readme"] = "dd-mm-yyyy" in z.read("README.txt").decode()
dump = admin.get("/admin/export/json").json()
rows = dump["tables"]["maintenance_payments"]["rows"]
out["json_iso"] = [bool(ISO.fullmatch(r["paid_on"])) for r in rows if r["paid_on"]][:3]
out["json_exported_at_iso"] = bool(re.match(r"\d{4}-\d{2}-\d{2}T", dump["exported_at"]))
# payee report exports
rc = admin.get("/admin/payee-report/export.csv?status=all").content.decode()
out["payee_csv"] = [leaks(rc)[:2], bool(re.search(r"\b\d\d-\d\d-\d{4}\b", rc))]
px = openpyxl.load_workbook(io.BytesIO(admin.get("/admin/payee-report/export.xlsx?status=all").content))
out["payee_xlsx"] = [t for t in sheet_text(px) if isinstance(t, str) and leaks(t)][:3]
out["payee_xlsx_generated"] = bool(re.search(r"Generated \d\d-\d\d-\d{4} \d\d:\d\d IST", px["Payments"]["B1"].value))
out["pdfs"] = [admin.get("/admin/payee-report/export.pdf?status=all").status_code, admin.get("/admin/export/pdf").status_code,
               owner.get(f"/owner/statement.pdf").status_code]
'''
    out = run_scratch(tmp_path, SCAN + code)
    assert out["xlsx_text_leaks"] == [] and out["xlsx_date_formats"] and all(f.startswith("dd-mm-yyyy") for f in out["xlsx_date_formats"])
    assert out["xlsx_about"] and ("IST" in str(out["xlsx_about"][0]))
    assert out["csv_bad"] == {}
    assert out["csv_paid_on"] and all(len(v) == 10 and v[2] == "-" and v[5] == "-" for v in out["csv_paid_on"])
    assert out["csv_readme"] is True
    assert out["json_iso"] and all(out["json_iso"]) and out["json_exported_at_iso"] is True
    assert out["payee_csv"][0] == [] and out["payee_csv"][1] is True
    assert out["payee_xlsx"] == [] and out["payee_xlsx_generated"] is True
    assert out["pdfs"] == [200, 200, 200]
