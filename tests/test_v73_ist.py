"""v7.3 (in place): every displayed time is IST (stored values stay UTC); "today" is the IST date; Pay now footer wording."""
import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from app import dates
from tests.test_v71_empty_and_review import run_scratch

ROOT = Path(__file__).resolve().parents[1]
FOOTER = ("Your payment is recorded right away against this due. The admin then checks the UTR against the UPI statement; "
          "if it can't be matched, the admin rejects it and the due comes back.")


def test_utc_to_ist_conversion_around_midnight():
    assert dates.format_datetime(datetime(2026, 10, 19, 18, 29, 59)) == "19-10-2026 23:59 IST"
    assert dates.format_datetime(datetime(2026, 10, 19, 18, 30)) == "20-10-2026 00:00 IST"  # next day in India
    assert dates.format_datetime(datetime(2026, 12, 31, 20, 0)) == "01-01-2027 01:30 IST"  # year rollover
    assert dates.format_datetime(datetime(2026, 10, 19, 11, 7, tzinfo=timezone.utc)) == "19-10-2026 16:37 IST"
    assert dates.format_datetime(None) == "—"
    assert dates.format_datetime(date(2026, 10, 5)) == "05-10-2026"  # a date-only value is never converted
    assert dates.format_date(datetime(2026, 10, 19, 18, 30)) == "19-10-2026"  # format_date of a datetime stays a plain date
    assert dates.to_ist(datetime(2026, 10, 19, 18, 30)).date() == date(2026, 10, 20)
    assert dates.ist_today(datetime(2026, 10, 19, 18, 29)) == date(2026, 10, 19)
    assert dates.ist_today(datetime(2026, 10, 19, 18, 30)) == date(2026, 10, 20)
    assert dates.ist_today(datetime(2026, 10, 20, 0, 0, tzinfo=timezone.utc)) == date(2026, 10, 20)
    # an IST calendar day starts 5h30 before UTC midnight
    assert dates.ist_day_start_utc(date(2026, 10, 20)) == datetime(2026, 10, 19, 18, 30)
    assert dates.display_dates_in_text('{"at": "2026-10-19T18:30:00"}') == '{"at": "20-10-2026 00:00 IST"}'
    assert dates.IST.utcoffset(None) == timedelta(hours=5, minutes=30)


def test_app_code_never_asks_the_server_clock_for_a_local_date():
    offenders = []
    for path in (ROOT / "app").rglob("*.py"):
        text = path.read_text()
        for pattern in (r"\bdate\.today\(", r"\bdatetime\.now\(\s*\)", r"\bdatetime\.today\("):
            if re.search(pattern, text):
                offenders.append((str(path.relative_to(ROOT)), pattern))
    assert offenders == []


CODE = r'''
import csv, io, re, zipfile
import openpyxl
from datetime import datetime
owner = login("asrk.sastry")
admin.post("/expenses", data={"description": "Bulbs", "category": "Common", "amount": "300", "expense_date": "2026-10-03"})
admin.post("/admin/recurring/tasks", data={"title": "Lift service", "day_of_month": "22", "due_mode": "by", "interval_months": "1", "notes": ""})
db.expire_all()
# pin every audit row to 2026-10-19 18:30 UTC = 20-10-2026 00:00 IST, one of them 18:29 UTC = 19-10-2026 23:59 IST
rows = db.scalars(select(AuditLog).order_by(AuditLog.id)).all()
out["audit_rows"] = len(rows)
for a in rows:
    a.created_at = datetime(2026, 10, 19, 18, 30)
rows[0].created_at = datetime(2026, 10, 19, 18, 29)
db.commit()
audit = admin.get("/audit").text
out["audit_page"] = ["20-10-2026 00:00 IST" in audit, "19-10-2026 23:59 IST" in audit, "19-10-2026 18:30" in audit]
act = admin.get("/admin/activity").text
out["activity_page"] = ["20-10-2026 00:00 IST" in act, "19-10-2026 18:30" in act]
dash = admin.get("/").text
out["dashboard_card"] = ["20-10-2026 00:00 IST" in dash, "19-10-2026 18:30" in dash]
# the activity date filter works on IST days: 20-10 includes the 00:00 IST rows, 19-10 only the 23:59 one
def count(df, dt):
    page = admin.get(f"/admin/activity?date_from={df}&date_to={dt}").text
    return page.count("20-10-2026 00:00 IST"), page.count("19-10-2026 23:59 IST")
out["filter_20"] = count("2026-10-20", "2026-10-20")
out["filter_19"] = count("2026-10-19", "2026-10-19")
# exports: CSV and Excel carry IST, the JSON backup stays UTC
z = zipfile.ZipFile(io.BytesIO(admin.get("/admin/export/csv").content))
logs = z.read("audit_logs.csv").decode()
out["csv"] = ["20-10-2026 00:00 IST" in logs, "19-10-2026 18:30" in logs, "UTC" in z.read("README.txt").decode(), "IST" in z.read("README.txt").decode()]
xl = openpyxl.load_workbook(io.BytesIO(admin.get("/admin/export/excel").content))
ws = xl["Audit logs"]
head = [c.value for c in ws[1]]
col = next(i for i, h in enumerate(head) if h and "created" in h.lower())
cells = [r[col] for r in ws.iter_rows(min_row=2)]
out["xlsx"] = [sorted({c.value.strftime("%d-%m-%Y %H:%M") for c in cells if isinstance(c.value, datetime)}), sorted({c.number_format for c in cells})]
dump = admin.get("/admin/export/json").json()
stored = {r["created_at"] for r in dump["tables"]["audit_logs"]["rows"]}
out["json"] = sorted(stored)
out["json_exported_at_utc"] = dump["exported_at"]
out["now_utc"] = datetime.utcnow().isoformat()
# the file names keep a stamp, now in IST
out["fname"] = admin.get("/admin/export/json").headers["content-disposition"]
out["ist_now"] = __import__("app.dates", fromlist=["x"]).ist_now().strftime("%Y%m%d-%H%M")
'''


def test_audit_activity_and_exports_show_ist_while_the_backup_stays_utc(tmp_path):
    out = run_scratch(tmp_path, CODE)
    assert out["audit_rows"] >= 2
    assert out["audit_page"] == [True, True, False]
    assert out["activity_page"] == [True, False]
    assert out["dashboard_card"] == [True, False]
    assert out["filter_20"][0] >= 1 and out["filter_20"][1] == 0
    assert out["filter_19"][0] == 0 and out["filter_19"][1] >= 1
    assert out["csv"] == [True, False, True, True]
    assert out["xlsx"][0] == ["19-10-2026 23:59", "20-10-2026 00:00"] and out["xlsx"][1] == ['dd-mm-yyyy hh:mm "IST"']
    assert out["json"] == ["2026-10-19T18:29:00", "2026-10-19T18:30:00"]
    stamp = re.search(r"dump-(\d{8})-(\d{6})\.json", out["fname"])
    assert stamp and stamp[1] == out["ist_now"][:8]  # file name stamp is IST
    exported = datetime.fromisoformat(out["json_exported_at_utc"])  # naive UTC, like every stored time
    assert abs((exported - datetime.fromisoformat(out["now_utc"])).total_seconds()) < 120


PAY = r'''
from app.dates import ist_today
owner = login("asrk.sastry")
sast = db.scalar(select(User).where(User.username == "asrk.sastry")); flat = sast.flats[0].id
period = db.scalar(select(LedgerPeriod).order_by(LedgerPeriod.year.desc(), LedgerPeriod.month.desc()))
mp = db.scalar(select(MaintenancePayment).where(MaintenancePayment.period_id == period.id, MaintenancePayment.flat_id == flat))
admin.post(f"/maintenance/{mp.id}/revert")
db.expire_all()
out["unpaid_before"] = not db.get(MaintenancePayment, mp.id).is_paid
sheet = owner.get(f"/?view=simple&period_id={period.id}&more=1&pay=m-{period.id}-{flat}").text
out["footer"] = "sv-note" in sheet and sheet.split('class="sv-note">')[1].split("</p>")[0]
out["old_wording"] = "nothing is marked paid" in sheet
r = owner.post("/owner/maintenance", data={"flat_id": flat, "period_id": period.id, "paid_on": ist_today().isoformat(), "amount": "",
                                           "payment_mode": "UPI", "payment_note": "UTR123456789012"}, follow_redirects=False)
db.expire_all()
row = db.get(MaintenancePayment, mp.id)
out["counted_at_once"] = [r.status_code, row.is_paid, row.submitted_by_user_id == sast.id]
subs = owner.get("/owner/submissions").text
out["submissions_copy"] = "counts as soon as you submit it" in subs and "the due comes back" in subs
admin.post(f"/admin/submissions/maintenance_payment/{mp.id}/reject", data={"reason": "UTR not on the statement", "back": "/admin/actions"})
db.expire_all()
out["after_reject_due_back"] = not db.get(MaintenancePayment, mp.id).is_paid
subs = owner.get("/owner/submissions").text
out["owner_sees_rejected"] = "Rejected" in subs and "UTR not on the statement" in subs
'''


def test_pay_now_footer_matches_what_actually_happens(tmp_path):
    out = run_scratch(tmp_path, PAY)
    assert out["unpaid_before"] is True
    assert out["footer"] == FOOTER and out["old_wording"] is False
    assert out["counted_at_once"] == [303, True, True]
    assert out["submissions_copy"] is True
    assert out["after_reject_due_back"] is True and out["owner_sees_rejected"] is True
