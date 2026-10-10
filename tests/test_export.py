import csv
import io
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from app.config import settings
from app.services import backup as backup_svc
from app.services import export as export_svc


def test_excel_export(admin_client):
    r = admin_client.get("/admin/export/excel")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "Expenses" in wb.sheetnames and "Audit logs" in wb.sheetnames
    ws = wb["Expenses"]
    headers = [c.value for c in ws[1]]
    assert "Expense date" in headers and "Settled?" in headers
    col = headers.index("Expense date") + 1
    assert ws.cell(row=2, column=col).number_format == "dd-mm-yyyy"
    user_headers = [c.value for c in wb["Users"][1]]
    assert "Password hash" not in user_headers


def test_csv_zip_export(admin_client):
    r = admin_client.get("/admin/export/csv")
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    names = set(zf.namelist())
    assert {"expenses.csv", "audit_logs.csv", "users.csv", "flats.csv"} <= names
    rows = list(csv.DictReader(io.StringIO(zf.read("expenses.csv").decode())))
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2}", rows[0]["expense_date"])
    assert "password_hash" not in zf.read("users.csv").decode().splitlines()[0]


def test_pdf_and_sqlite_downloads(admin_client, seeded_db):
    from app.models import LedgerPeriod

    oct_id = seeded_db.query(LedgerPeriod).filter_by(year=2026, month=10).one().id
    for q in ("", "?period_id=all", f"?period_id={oct_id}"):
        r = admin_client.get(f"/admin/export/pdf{q}")
        assert r.status_code == 200 and r.content.startswith(b"%PDF")
    r = admin_client.get("/admin/export/sqlite")
    assert r.status_code == 200 and r.content.startswith(b"SQLite format 3")


def test_json_dump_has_every_table(admin_client, seeded_db):
    data = admin_client.get("/admin/export/json").json()
    assert data["format"] == export_svc.DUMP_FORMAT
    assert "schema_version" in data
    counts = export_svc.row_counts(seeded_db)
    assert set(data["tables"]) == set(counts)
    assert data["row_counts"] == counts
    assert data["row_counts"]["audit_logs"] > 0


def test_export_wipe_restore_round_trip(seeded_db):
    before = export_svc.row_counts(seeded_db)
    dump = json.loads(export_svc.dump_json_bytes(seeded_db))

    for t in reversed(export_svc.tables()):
        seeded_db.execute(t.delete())
    seeded_db.commit()
    assert sum(export_svc.row_counts(seeded_db).values()) == 0

    export_svc.restore_dump(seeded_db, dump)
    seeded_db.commit()
    assert export_svc.row_counts(seeded_db) == before
    again = json.loads(export_svc.dump_json_bytes(seeded_db))
    assert again["tables"] == dump["tables"]


def test_restore_via_admin_page_with_confirmation(admin_client, seeded_db):
    dump_bytes = admin_client.get("/admin/export/json").content
    before = export_svc.row_counts(seeded_db)
    seeded_db.close()

    r = admin_client.post("/admin/export/restore", files={"dump": ("dump.json", dump_bytes, "application/json")})
    assert r.status_code == 200 and "Confirm restore" in r.text
    token = re.search(r'name="token" value="([^"]+)"', r.text).group(1)

    # Wrong confirmation word cancels.
    r = admin_client.post("/admin/export/restore/confirm", data={"token": token, "confirm_text": "yes"})
    assert r.status_code == 400 and "cancelled" in r.text

    r = admin_client.post("/admin/export/restore", files={"dump": ("dump.json", dump_bytes, "application/json")})
    token = re.search(r'name="token" value="([^"]+)"', r.text).group(1)
    backups_before = set(backup_svc.list_backups())
    r = admin_client.post("/admin/export/restore/confirm", data={"token": token, "confirm_text": "RESTORE"})
    assert r.status_code == 200 and "Restore complete" in r.text
    new_backups = set(backup_svc.list_backups()) - backups_before
    assert any("pre-restore" in p.name for p in new_backups)

    from app.database import SessionLocal

    db = SessionLocal()
    after = export_svc.row_counts(db)
    db.close()
    assert after["audit_logs"] == before["audit_logs"] + 1  # the restore itself is audited
    assert {k: v for k, v in after.items() if k != "audit_logs"} == {
        k: v for k, v in before.items() if k != "audit_logs"
    }
    assert admin_client.get("/admin/users").status_code == 200


def test_restore_rejects_garbage(admin_client):
    r = admin_client.post("/admin/export/restore", files={"dump": ("x.json", b'{"hello": 1}', "application/json")})
    assert r.status_code == 400 and "Not a Nest5" in r.text
    r = admin_client.post("/admin/export/restore", files={"dump": ("x.json", b"not json", "application/json")})
    assert r.status_code == 400


def test_startup_backup_keeps_last_30(seeded_db):
    bdir = Path(settings.backup_dir)
    bdir.mkdir(parents=True, exist_ok=True)
    for i in range(35):
        (bdir / f"nest5-20200101-0000{i:02d}.db").write_bytes(b"x")
    dest = backup_svc.make_backup()
    assert dest is not None and dest.exists()
    assert re.fullmatch(r"nest5-\d{8}-\d{6}\.db", dest.name)
    files = sorted(bdir.glob("nest5-*.db"))
    assert len(files) == 30
    assert dest in files
