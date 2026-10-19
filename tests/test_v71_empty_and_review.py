"""v7.1: admin empty states, Actions board (Mark reviewed, tidy list), owner payment rejection."""
import json
import os
import re
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from sqlalchemy import select

from app.models import (
    AuditLog,
    CancelRequest,
    CorpusCampaign,
    CorpusPayment,
    Expense,
    MaintenancePayment,
    Project,
    ProjectContribution,
    ProjectPayment,
    SubmissionRejection,
    SubmissionReview,
)
from app.services import actions as actions_svc
from app.services import corpus_campaigns as campaigns_svc
from app.services import maintenance_fee as fee_svc
from app.services import reminders as reminders_svc
from app.services import submission_review as review_svc
from tests.conftest import OWNER_PASSWORD, make_client
from tests.test_cancel_everywhere import (
    PAGES,
    _cancel,
    _flats,
    _free_maintenance_row,
    _oct,
    _pages_without,
    _restore_maintenance_row,
    totals,
)

ROOT = Path(__file__).resolve().parents[1]


def strip_history(page: str) -> str:
    """Recent activity is history and may name rejected / cancelled entries; every other area may not."""
    return re.sub(r'<details class="card dashboard-activity.*?</details>', "", page, flags=re.S)


def run_scratch(tmp_path, code: str, *, prelude: str = "") -> str:
    """Run `code` against a brand-new seeded database (own files, never the shared test DB)."""
    env = {**os.environ, "DATABASE_URL": f"sqlite:///{tmp_path / 'scratch.db'}",
           "INSTANCE_DIR": str(tmp_path / "i"), "BACKUP_DIR": str(tmp_path / "b")}
    seed = subprocess.run([sys.executable, "scripts/seed.py"], cwd=ROOT, env=env, capture_output=True, text=True)
    assert seed.returncode == 0, seed.stderr
    head = (
        "import json\n"
        "from fastapi.testclient import TestClient\n"
        "from sqlalchemy import select, text\n"
        "from app.main import app\n"
        "from app.database import SessionLocal\n"
        "from app.models import *\n"
        "from app.services import users as u\n"
        "db = SessionLocal(); u.ensure_bootstrap_users(db); db.commit()\n"
        "for usr in db.scalars(select(User)).all():\n"
        "    u.set_password(usr, 'pw-pw-pw-1', must_change=False)\n"
        "db.commit()\n"
        "def login(name):\n"
        "    c = TestClient(app)\n"
        "    assert c.post('/login', data={'username': name, 'password': 'pw-pw-pw-1'}, follow_redirects=False).status_code == 303\n"
        "    return c\n"
        "admin = login('admin')\n"
        "out = {}\n"
    )
    ran = subprocess.run([sys.executable, "-c", head + prelude + code + "\nprint('JSON=' + json.dumps(out))"],
                         cwd=ROOT, env=env, capture_output=True, text=True)
    assert "JSON=" in ran.stdout, ran.stderr[-2500:]
    return json.loads(ran.stdout.split("JSON=", 1)[1])


# ---- 1. Admin empty states -----------------------------------------------------------------


EMPTY_SCRIPT = """
from datetime import datetime
for e in db.scalars(select(Expense)).all():
    e.deleted_at = datetime.utcnow()
db.commit()
pages = {}
for url in ('/', '/corpus', '/projects', '/expenses', '/payments/corpus', '/payments/projects', '/payments/expenses',
            '/corpus?add=goal', '/projects?add=project', '/expenses?add=expense', '/flats/1', '/admin/actions'):
    r = admin.get(url)
    assert r.status_code == 200, (url, r.status_code)
    pages[url] = r.text
out['pages'] = pages
owner = login('asrk.sastry')
out['owner_corpus'] = owner.get('/corpus').text
out['owner_projects'] = owner.get('/projects').text
"""


def test_admin_sees_no_empty_cards_without_goals_projects_expenses(tmp_path):
    pages = run_scratch(tmp_path, EMPTY_SCRIPT)["pages"]
    dash = pages["/"]
    for needle in ("+ Add corpus goal", "+ Add project", "+ Add expense"):
        assert needle in dash, needle
    assert "<th>Corpus</th>" not in dash and "Corpus</dt>" not in dash
    assert "settlement-card" not in dash and "Unsettled owner expenses" not in dash
    assert "Paid from Account Balance" not in dash and "Available to settle now" not in dash
    body = re.sub(r"<nav.*?</nav>", "", dash, flags=re.S)
    assert 'href="/corpus"' not in body and 'href="/projects"' not in body
    assert "<th>Maintenance" in dash  # maintenance always exists

    corpus = pages["/corpus"]
    assert "+ Add corpus goal" in corpus and 'action="/corpus/campaigns"' in corpus  # create form still there
    for gone in ("Per-flat status", "No corpus goals", "Target (all goals)", "Pending (open goals)",
                 "Record a corpus payment", "<h3 style=\"margin-top:0\">Goals</h3>", "one-off"):
        assert gone not in corpus, gone
    assert "<details id=\"add-goal\" class=\"add-details no-print\">" in corpus
    assert "<details id=\"add-goal\" class=\"add-details no-print\" open>" in pages["/corpus?add=goal"]

    projects = pages["/projects"]
    assert "+ Add project" in projects and 'action="/projects"' in projects
    for gone in ("No projects yet", "<thead>", "Total project cost", "Project costs are split"):
        assert gone not in projects, gone
    assert "<details id=\"add-project\" class=\"add-details no-print\" open>" in pages["/projects?add=project"]

    expenses = pages["/expenses"]
    assert "+ Add expense" in expenses and 'action="/expenses"' in expenses
    for gone in ("No expenses for this month", "<thead>", "Owner advances to reimburse", "Account Balance available"):
        assert gone not in expenses, gone
    assert "<details id=\"add-expense\" class=\"add-details no-print\" open>" in pages["/expenses?add=expense"]

    for url, link in (("/payments/corpus", "/corpus?add=goal#add-goal"), ("/payments/projects", "/projects?add=project#add-project"),
                      ("/payments/expenses", "/expenses?add=expense#add-expense")):
        page = pages[url]
        assert link in page and "<table" not in page and "No " not in page.split("</nav>", 1)[1].split("<script")[0], url

    flat = pages["/flats/1"]
    assert "You owe (projects)" not in flat and "You owe (corpus)" not in flat
    assert "No advances recorded" not in flat and "No open corpus collections" not in flat

    actions = pages["/admin/actions"]
    assert "All clear" in actions and "Nothing to do" not in actions


def test_owner_views_unchanged_when_empty(tmp_path):
    data = run_scratch(tmp_path, EMPTY_SCRIPT)
    assert "+ Add corpus goal" not in data["owner_corpus"] and 'action="/corpus/campaigns"' not in data["owner_corpus"]
    assert "+ Add project" not in data["owner_projects"]


def test_empty_states_disappear_once_data_exists(tmp_path):
    code = EMPTY_SCRIPT + """
admin.post('/corpus/campaigns', data={'name': 'Lift fund', 'target_amount': '50000'})
admin.post('/projects', data={'name': 'Terrace', 'total_cost': '10000'})
from app.services import periods as periods_svc
pid = periods_svc.resolve_period(db, None).id
admin.post('/expenses', data={'expense_date': '2026-10-05', 'category': 'Common', 'description': 'Bulbs', 'amount': '640', 'period_id': str(pid)})
out['after'] = {u: admin.get(u).text for u in ('/', '/corpus', '/projects', '/expenses')}
"""
    after = run_scratch(tmp_path, code)["after"]
    assert "Per-flat status" in after["/corpus"] and "Lift fund" in after["/corpus"] and 'id="add-goal"' in after["/corpus"]
    assert "<details id=\"add-goal\"" not in after["/corpus"]  # form is a normal card now
    assert "Terrace" in after["/projects"] and "<thead>" in after["/projects"]
    assert "Bulbs" in after["/expenses"] and "<thead>" in after["/expenses"]
    dash = after["/"]
    assert "<th>Corpus</th>" in dash and "Terrace" in dash and "Paid from Account Balance" in dash
    assert "Not set up yet" not in dash


def test_general_one_off_form_is_gone_but_general_goal_stays_internal(seeded_db, admin_client):
    general = campaigns_svc.general_campaign(seeded_db, create=True)
    seeded_db.commit()
    gid = general.id
    try:
        campaigns_svc.add_to_general_share(seeded_db, _flats(seeded_db)["201"], 500)
        seeded_db.commit()
        for url in ("/corpus", "/payments/corpus"):
            page = admin_client.get(url).text
            assert "one-off" not in page.lower() and "Add to share" not in page, url
            assert 'method="post" action="/corpus"' not in page, url
        # General corpus is still a normal, deletable goal (v9.1 behaviour).
        r = admin_client.post(f"/corpus/goals/{gid}/delete", data={"confirm": "DELETE", "ack": "1"}, follow_redirects=False)
        assert r.status_code == 303
    finally:
        seeded_db.expire_all()
        g = seeded_db.get(CorpusCampaign, gid)
        if g is not None:
            for p in seeded_db.scalars(select(CorpusPayment).where(CorpusPayment.campaign_id == gid)).all():
                seeded_db.delete(p)
            seeded_db.delete(g)
            seeded_db.commit()
    assert admin_client.get("/corpus").status_code == 200


# ---- helpers for submissions ---------------------------------------------------------------


def _siva(db, set_password):
    set_password(db, "asrk.sastry", OWNER_PASSWORD)
    return make_client("asrk.sastry", OWNER_PASSWORD)


def _owner_project_payment(owner, db, project_name, flat_id, amount=100):
    db.expire_all()
    project = db.scalar(select(Project).where(Project.name == project_name))
    pp = db.scalar(select(ProjectPayment).where(ProjectPayment.project_id == project.id, ProjectPayment.flat_id == flat_id))
    r = owner.post("/owner/projects/pay", data={"share_id": pp.id, "amount": amount, "paid_on": date.today().isoformat(),
                                                "payment_mode": "Cash"}, follow_redirects=False)
    assert "owner_msg" in r.headers["location"], r.headers["location"]
    db.expire_all()
    return db.scalars(select(ProjectContribution).where(ProjectContribution.project_payment_id == pp.id)).all()[-1]


def _drop_project(db, project_id):
    db.expire_all()
    for pc in db.scalars(select(ProjectContribution).where(ProjectContribution.project_id == project_id)).all():
        db.delete(pc)
    db.flush()
    for pp in db.scalars(select(ProjectPayment).where(ProjectPayment.project_id == project_id)).all():
        db.delete(pp)
    p = db.get(Project, project_id)
    if p is not None:
        db.delete(p)
    db.commit()


def _clear_reviews(db):
    for row in db.scalars(select(SubmissionReview)).all():
        db.delete(row)
    for row in db.scalars(select(SubmissionRejection)).all():
        db.delete(row)
    db.commit()


# ---- 2. Actions board: Mark reviewed ---------------------------------------------------------


def test_mark_reviewed_one_then_all_persists_and_keeps_history(seeded_db, admin_client, set_password):
    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    admin_client.post("/projects", data={"name": "Review project", "total_cost": "5000"})
    seeded_db.expire_all()
    project = seeded_db.scalar(select(Project).where(Project.name == "Review project"))
    pid = project.id
    try:
        c1 = _owner_project_payment(owner, seeded_db, "Review project", flats["201"], 111)
        c2 = _owner_project_payment(owner, seeded_db, "Review project", flats["201"], 222)
        pending = review_svc.pending(seeded_db)
        mine = [p for p in pending if p["record_id"] in (c1.id, c2.id) and p["kind"] == "project_contribution"]
        assert len(mine) == 2
        page = admin_client.get("/admin/actions").text
        assert "Mark all reviewed" in page and page.count("Mark reviewed") >= 2
        assert "₹111" in page and "₹222" in page and "Project Review project" in page and "Flat 201" in page
        assert f'href="/projects/{pid}"' in page
        dash_total = actions_svc.build(seeded_db)["total"]

        first = next(p for p in mine if p["record_id"] == c1.id)
        r = admin_client.post(f"/admin/actions/submissions/{first['audit_id']}/reviewed", data={"back": "/admin/actions"},
                              follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"].startswith("/admin/actions?msg=")
        seeded_db.expire_all()
        row = seeded_db.scalar(select(SubmissionReview).where(SubmissionReview.audit_id == first["audit_id"]))
        assert row.kind == "project_contribution" and row.record_id == c1.id and row.reviewed_by == "admin"
        assert row.reviewed_at is not None
        page = admin_client.get("/admin/actions").text
        assert "₹111" not in page and "₹222" in page
        assert actions_svc.build(seeded_db)["total"] == dash_total - 1
        # the dashboard Actions badge follows
        assert f'Actions <span class="badge warn">{dash_total - 1}</span>' in admin_client.get("/").text
        # marking the same one again is harmless
        again = admin_client.post(f"/admin/actions/submissions/{first['audit_id']}/reviewed", follow_redirects=False)
        assert again.status_code == 303 and "error=" not in again.headers["location"]
        assert len(seeded_db.scalars(select(SubmissionReview).where(SubmissionReview.audit_id == first["audit_id"])).all()) == 1
        # persists across a new login
        fresh = make_client("admin", "nest5admin")
        assert "₹111" not in fresh.get("/admin/actions").text

        # history keeps it: payment still counts, audit log + Activity still list it
        assert seeded_db.get(ProjectContribution, c1.id) is not None
        assert "₹111" in admin_client.get("/admin/activity").text
        assert "submitted project payment ₹111" in admin_client.get("/audit").text

        r = admin_client.post("/admin/actions/submissions-reviewed", follow_redirects=False)
        assert r.headers["location"].startswith("/admin/actions?msg=")
        seeded_db.expire_all()
        assert not [p for p in review_svc.pending(seeded_db) if p["record_id"] in (c1.id, c2.id)]
        assert "₹222" not in admin_client.get("/admin/actions").text
        assert "Project Review project" not in strip_history(admin_client.get("/").text)
        actions = {a.action for a in seeded_db.scalars(select(AuditLog).where(AuditLog.entity_type == "submission")).all()}
        assert {"review", "review_all"} <= actions

        # a new submission after "mark all" shows up again, un-reviewed
        _owner_project_payment(owner, seeded_db, "Review project", flats["201"], 333)
        assert "₹333" in admin_client.get("/admin/actions").text
    finally:
        _drop_project(seeded_db, pid)
        _clear_reviews(seeded_db)
        for log in seeded_db.scalars(select(AuditLog).where(AuditLog.entity_type == "submission")).all():
            pass


def test_owner_cannot_mark_reviewed_or_reject(seeded_db, owner_client):
    for url in ("/admin/actions/submissions/1/reviewed", "/admin/actions/submissions-reviewed",
                "/admin/submissions/maintenance_payment/1/reject"):
        assert owner_client.post(url).status_code == 403, url
    assert owner_client.get("/admin/actions").status_code == 403


def test_legacy_mark_all_timestamp_is_still_honoured(seeded_db, admin_client, set_password):
    """Submissions listed at upgrade time stay pending; only those before the old 'since' stamp are done."""
    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    admin_client.post("/projects", data={"name": "Legacy project", "total_cost": "5000"})
    seeded_db.expire_all()
    pid = seeded_db.scalar(select(Project).where(Project.name == "Legacy project")).id
    from app.services import app_settings as settings_svc
    try:
        c = _owner_project_payment(owner, seeded_db, "Legacy project", flats["201"], 444)
        assert any(p["record_id"] == c.id for p in review_svc.pending(seeded_db))  # nothing auto-cleared
        settings_svc.set_value(seeded_db, review_svc.LEGACY_REVIEWED_KEY,
                               (datetime.utcnow() + timedelta(minutes=1)).isoformat(), updated_by="admin")
        seeded_db.commit()
        assert not any(p["record_id"] == c.id for p in review_svc.pending(seeded_db))
    finally:
        from app.models import AppSetting
        row = seeded_db.get(AppSetting, review_svc.LEGACY_REVIEWED_KEY)
        if row is not None:
            seeded_db.delete(row)
            seeded_db.commit()
        _drop_project(seeded_db, pid)
        _clear_reviews(seeded_db)


def test_actions_board_is_grouped_one_line_per_item_and_hides_empty_groups(seeded_db, admin_client, set_password):
    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    admin_client.post("/projects", data={"name": "Board project", "total_cost": "5000"})
    seeded_db.expire_all()
    pid = seeded_db.scalar(select(Project).where(Project.name == "Board project")).id
    try:
        c = _owner_project_payment(owner, seeded_db, "Board project", flats["201"], 250)
        page = admin_client.get("/admin/actions").text
        assert "Needs your review" in page and "New owner submissions to review" in page
        row = re.search(r'<li class="actions-row">(?:(?!</li>).)*₹250(?:(?!</li>).)*</li>', page, flags=re.S).group(0)
        for needle in ("Project Board project", "Flat 201", "₹250", "by asrk.sastry", f'href="/projects/{pid}"',
                       "Mark reviewed", "Reject"):
            assert needle in row, needle
        assert row.count("<li") == 1
        for title in ("Overdue corpus", "Overdue asset tasks", "Overdue project shares"):
            assert title not in page
        assert "Nothing to do" not in page and "No new owner submissions" not in page
        assert page.count("New owner submissions to review") == 1
        dash = admin_client.get("/").text
        assert "Mark reviewed" in dash and "Reject" not in dash.split("dashboard-actions", 1)[1].split("</section>")[0]
    finally:
        _drop_project(seeded_db, pid)
        _clear_reviews(seeded_db)


def test_actions_board_all_clear_message(tmp_path):
    out = run_scratch(tmp_path, "out['page'] = admin.get('/admin/actions').text\nout['dash'] = admin.get('/').text")
    assert "All clear" in out["page"]
    assert "Open items" not in out["page"]
    assert "dashboard-actions" not in out["dash"]


# ---- 3. Reject owner payments: gone from every list / total, due again --------------------------


def _reject(admin, kind, record_id, reason="", back="/admin/actions"):
    return admin.post(f"/admin/submissions/{kind}/{record_id}/reject", data={"reason": reason, "back": back},
                      follow_redirects=False)


def test_reject_owner_maintenance_payment_restores_the_due(seeded_db, admin_client, set_password):
    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    mp, saved = _free_maintenance_row(seeded_db, flats["201"])
    mp_id, period_id = mp.id, mp.period_id
    try:
        before = totals(seeded_db, [flats["201"], flats["101"]])
        owner.post("/owner/maintenance", data={"flat_id": flats["201"], "period_id": period_id, "paid_on": "2026-10-06",
                                               "amount": "2345", "payment_mode": "UPI", "payment_note": "UTR987654321"})
        during = totals(seeded_db, [flats["201"], flats["101"]])
        assert during != before and during["fund_balance"] == before["fund_balance"] + 2345
        assert "₹2345" in admin_client.get("/admin/actions").text
        assert "Reject" in admin_client.get(f"/payments/maintenance?period_id={period_id}").text

        r = _reject(admin_client, "maintenance_payment", mp_id, "Not in the bank statement",
                    back=f"/payments/maintenance?period_id={period_id}")
        assert r.status_code == 303 and "msg=" in r.headers["location"] and r.headers["location"].count("?") == 1
        seeded_db.expire_all()
        mp = seeded_db.get(MaintenancePayment, mp_id)
        assert not mp.is_paid and mp.paid_on is None and mp.submitted_by_user_id is None
        assert mp.amount == fee_svc.fee_for_period(seeded_db, mp.period)  # the month's fee is due again, not ₹2345

        assert totals(seeded_db, [flats["201"], flats["101"]]) == before
        _pages_without([admin_client, owner], "₹2345",
                       extra=[f"/flats/{flats['201']}", f"/statements/{period_id}", "/?view=simple"])
        due = [i for i in reminders_svc.reminders_for_flats(seeded_db, [flats["201"]])["items"] if i["kind"] == "Maintenance"]
        assert any("October" in i["label"] for i in due)
        simple = owner.get("/?view=simple").text
        total_due = reminders_svc.reminders_for_flats(seeded_db, [flats["201"]])["total"]
        assert total_due > 0 and f"<span>You owe</span><strong>₹{total_due}</strong>" in simple
        assert "Your dues to pay" in simple and "sv-due-row" in simple

        # kept only in the submitting owner's My submissions, as Rejected, with the reason
        mine = owner.get("/owner/submissions").text
        assert "Rejected" in mine and "Not in the bank statement" in mine and "2345" in mine
        assert "Not in the bank statement" not in strip_history(admin_client.get("/").text)
        rej = seeded_db.scalar(select(SubmissionRejection).where(SubmissionRejection.record_id == mp_id))
        assert rej.record_type == "maintenance_payment" and rej.reason == "Not in the bank statement"
        assert rej.rejected_by == "admin" and rej.amount == 2345
        entry = seeded_db.scalar(select(AuditLog).where(AuditLog.entity_type == "maintenance_payment",
                                                         AuditLog.entity_id == mp_id, AuditLog.action == "reject"))
        assert entry is not None and entry.user_label == "admin"
        assert not any(p["record_id"] == mp_id for p in review_svc.pending(seeded_db))

        # rejecting twice / an admin-recorded payment is refused
        again = _reject(admin_client, "maintenance_payment", mp_id)
        assert "error=" in again.headers["location"]
    finally:
        _restore_maintenance_row(seeded_db, mp_id, saved)
        _clear_reviews(seeded_db)


def test_reject_owner_corpus_payment_updates_goal_progress(seeded_db, admin_client, set_password):
    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    admin_client.post("/corpus/campaigns", data={"name": "Reject drive", "target_amount": "50000"})
    seeded_db.expire_all()
    camp = seeded_db.scalar(select(CorpusCampaign).where(CorpusCampaign.name == "Reject drive"))
    cid = camp.id
    try:
        before = totals(seeded_db, [flats["201"]])
        owner.post("/corpus/pay", data={"target": f"{flats['201']}:{cid}", "amount": "1234", "paid_on": date.today().isoformat(),
                                        "payment_mode": "Cash"})
        seeded_db.expire_all()
        cp = seeded_db.scalar(select(CorpusPayment).where(CorpusPayment.campaign_id == cid))
        assert campaigns_svc.campaign_summary(seeded_db, camp)["paid_total"] == 1234
        assert "Reject" in admin_client.get("/payments/corpus").text
        cp_id = cp.id
        r = _reject(admin_client, "corpus_payment", cp_id, "Wrong flat", back="/payments/corpus")
        assert "msg=" in r.headers["location"]
        seeded_db.expire_all()
        assert seeded_db.get(CorpusPayment, cp_id) is None
        summary = campaigns_svc.campaign_summary(seeded_db, seeded_db.get(CorpusCampaign, cid))
        assert summary["paid_total"] == 0 and summary["progress_pct"] == 0
        assert totals(seeded_db, [flats["201"]]) == before
        _pages_without([admin_client, owner], "₹1234", extra=[f"/flats/{flats['201']}", f"/corpus/goals/{cid}", "/?view=simple"])
        assert "0%" in admin_client.get(f"/corpus/goals/{cid}").text
        mine = owner.get("/owner/submissions").text
        assert "Rejected" in mine and "Wrong flat" in mine
        # the share is due again for the owner
        assert any("Reject drive" in i["label"] for i in reminders_svc.reminders_for_flats(seeded_db, [flats["201"]])["items"])
    finally:
        seeded_db.expire_all()
        for p in seeded_db.scalars(select(CorpusPayment).where(CorpusPayment.campaign_id == cid)).all():
            seeded_db.delete(p)
        c = seeded_db.get(CorpusCampaign, cid)
        for sh in list(c.shares):
            seeded_db.delete(sh)
        seeded_db.delete(c)
        seeded_db.commit()
        _clear_reviews(seeded_db)


def test_reject_owner_project_payment_updates_paid_percent(seeded_db, admin_client, set_password):
    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    admin_client.post("/projects", data={"name": "Reject project", "total_cost": "5000"})
    seeded_db.expire_all()
    pid = seeded_db.scalar(select(Project).where(Project.name == "Reject project")).id
    try:
        before = totals(seeded_db, [flats["201"]])
        contrib = _owner_project_payment(owner, seeded_db, "Reject project", flats["201"], 650)
        assert "13%" in admin_client.get("/projects").text
        assert "Reject" in admin_client.get("/payments/projects").text
        contrib_id = contrib.id
        r = _reject(admin_client, "project_contribution", contrib_id, "Duplicate", back="/payments/projects")
        assert "msg=" in r.headers["location"]
        seeded_db.expire_all()
        assert seeded_db.get(ProjectContribution, contrib_id) is None
        share = seeded_db.scalar(select(ProjectPayment).where(ProjectPayment.project_id == pid, ProjectPayment.flat_id == flats["201"]))
        assert share.paid == 0 and share.status == "pending" and share.remaining == share.amount
        assert totals(seeded_db, [flats["201"]]) == before
        projects_page = admin_client.get("/projects").text
        assert "0%" in projects_page and "13%" not in projects_page
        _pages_without([admin_client, owner], "₹650", extra=[f"/flats/{flats['201']}", f"/projects/{pid}", "/?view=simple"])
        assert "Rejected" in owner.get("/owner/submissions").text and "Duplicate" in owner.get("/owner/submissions").text
        assert any("Reject project" in i["label"] for i in reminders_svc.reminders_for_flats(seeded_db, [flats["201"]])["items"])
    finally:
        _drop_project(seeded_db, pid)
        _clear_reviews(seeded_db)


def test_reject_owner_expense_uses_the_same_endpoint_and_admin_entries_cannot_be_rejected(seeded_db, admin_client, set_password):
    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    before = totals(seeded_db, [flats["201"]])
    owner.post("/owner/expenses", data={"flat_id": flats["201"], "description": "Reject bulbs", "category": "Common",
                                        "amount": "640", "expense_date": "2026-10-04"})
    seeded_db.expire_all()
    exp = seeded_db.scalar(select(Expense).where(Expense.description == "Reject bulbs"))
    try:
        assert "₹640" in admin_client.get("/admin/actions").text
        r = _reject(admin_client, "expense", exp.id, "Personal")
        assert "msg=" in r.headers["location"]
        seeded_db.expire_all()
        assert seeded_db.get(Expense, exp.id).status == "rejected"
        assert totals(seeded_db, [flats["201"]]) == before
        assert "Reject bulbs" not in admin_client.get("/admin/actions").text
        assert "Rejected" in owner.get("/owner/submissions").text

        # an admin-recorded maintenance payment is not an owner submission
        mp = seeded_db.scalar(select(MaintenancePayment).where(MaintenancePayment.is_paid.is_(True)))
        assert "error=" in _reject(admin_client, "maintenance_payment", mp.id).headers["location"]
        assert "error=" in _reject(admin_client, "nonsense", 1).headers["location"]
        assert "error=" in _reject(admin_client, "corpus_payment", 999999).headers["location"]
    finally:
        seeded_db.delete(seeded_db.get(Expense, exp.id))
        seeded_db.commit()
        _clear_reviews(seeded_db)


def test_reject_closes_a_pending_cancel_request_without_double_listing(seeded_db, admin_client, set_password):
    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    admin_client.post("/projects", data={"name": "Cancel then reject", "total_cost": "5000"})
    seeded_db.expire_all()
    pid = seeded_db.scalar(select(Project).where(Project.name == "Cancel then reject")).id
    try:
        contrib = _owner_project_payment(owner, seeded_db, "Cancel then reject", flats["201"], 500)
        contrib_id = contrib.id
        owner.post("/owner/cancel-requests", data={"record_type": "project_contribution", "record_id": contrib_id})
        assert "Cancel requests waiting" in admin_client.get("/admin/actions").text
        _reject(admin_client, "project_contribution", contrib_id, "no")
        seeded_db.expire_all()
        req = seeded_db.scalar(select(CancelRequest).where(CancelRequest.record_id == contrib_id,
                                                           CancelRequest.record_type == "project_contribution"))
        assert req.status == "approved" and "rejected" in req.decision_note.lower()
        assert "Cancel then reject" not in admin_client.get("/admin/actions").text
        mine = owner.get("/owner/submissions").text
        assert mine.count("Cancel then reject") == 1 and "Rejected" in mine
    finally:
        _drop_project(seeded_db, pid)
        _clear_reviews(seeded_db)
        for r in seeded_db.scalars(select(CancelRequest).where(CancelRequest.summary.like("%Cancel then reject%"))).all():
            seeded_db.delete(r)
        seeded_db.commit()


def test_exports_leave_out_cancelled_and_rejected_but_json_keeps_them(seeded_db, admin_client, set_password):
    import io
    import zipfile

    from openpyxl import load_workbook

    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    admin_client.post("/projects", data={"name": "Export project", "total_cost": "5000"})
    seeded_db.expire_all()
    pid = seeded_db.scalar(select(Project).where(Project.name == "Export project")).id
    owner.post("/owner/expenses", data={"flat_id": flats["201"], "description": "Cancelled export bill", "category": "Common",
                                        "amount": "4321", "expense_date": "2026-10-04"})
    seeded_db.expire_all()
    exp = seeded_db.scalar(select(Expense).where(Expense.description == "Cancelled export bill"))
    try:
        contrib = _owner_project_payment(owner, seeded_db, "Export project", flats["201"], 999)
        _cancel(owner, admin_client, seeded_db, "expense", exp.id)
        _reject(admin_client, "project_contribution", contrib.id, "Bad ref")
        xlsx = load_workbook(io.BytesIO(admin_client.get("/admin/export/excel").content))
        sheets = [ws for ws in xlsx.worksheets if ws.title not in ("Audit logs", "Cancel requests")]  # history sheets keep their lines
        text = " ".join(str(c.value) for ws in sheets for row in ws.iter_rows() for c in row if c.value is not None)
        assert "Cancelled export bill" not in text and "Bad ref" not in text
        z = zipfile.ZipFile(io.BytesIO(admin_client.get("/admin/export/csv").content))
        assert "Cancelled export bill" not in z.read("expenses.csv").decode()
        assert z.read("submission_rejections.csv").decode().strip().count("\n") == 0  # header only
        pdf = admin_client.get("/admin/export/pdf?period_id=")
        assert pdf.status_code == 200
        dump = admin_client.get("/admin/export/json").json()
        assert any(r["description"] == "Cancelled export bill" and r["deleted_at"] for r in dump["tables"]["expenses"]["rows"])
        assert any(r["reason"] == "Bad ref" for r in dump["tables"]["submission_rejections"]["rows"])
        assert "status" in dump["tables"]["expenses"]["columns"]
        assert "submission_reviews" in dump["tables"]
    finally:
        _drop_project(seeded_db, pid)
        seeded_db.delete(seeded_db.get(Expense, exp.id))
        seeded_db.commit()
        _clear_reviews(seeded_db)
        for r in seeded_db.scalars(select(CancelRequest).where(CancelRequest.summary.like("%Cancelled export bill%"))).all():
            seeded_db.delete(r)
        seeded_db.commit()


def test_cancel_approval_restores_simple_view_dues_and_progress(seeded_db, admin_client, set_password):
    owner = _siva(seeded_db, set_password)
    flats = _flats(seeded_db)
    admin_client.post("/projects", data={"name": "Cancel progress", "total_cost": "5000"})
    seeded_db.expire_all()
    pid = seeded_db.scalar(select(Project).where(Project.name == "Cancel progress")).id
    try:
        before = totals(seeded_db, [flats["201"]])
        contrib = _owner_project_payment(owner, seeded_db, "Cancel progress", flats["201"], 650)
        assert "13%" in admin_client.get("/projects").text
        _cancel(owner, admin_client, seeded_db, "project_contribution", contrib.id)
        assert totals(seeded_db, [flats["201"]]) == before
        page = admin_client.get("/projects").text
        assert "0%" in page and "13%" not in page
        simple = owner.get("/?view=simple").text
        assert "Cancel progress" in simple  # the full share is a due again
        _pages_without([admin_client, owner], "₹650", extra=[f"/projects/{pid}", f"/flats/{flats['201']}"])
    finally:
        _drop_project(seeded_db, pid)
        _clear_reviews(seeded_db)
        for r in seeded_db.scalars(select(CancelRequest).where(CancelRequest.summary.like("%Cancel progress%"))).all():
            seeded_db.delete(r)
        seeded_db.commit()
