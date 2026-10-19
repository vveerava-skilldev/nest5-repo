"""v7.3 (in place): admin-only Payee / UPI audit report: filters, totals, exports, access, read-only."""
from app.dates import ist_today
from tests.test_v71_empty_and_review import run_scratch
from tests.test_v73_payment_settings import SETUP as PAY_SETUP

OLD = "9940255076@ybl"
NEW = "new.society@oksbi"

SETUP = PAY_SETUP + r'''
import csv, io, re
import openpyxl
from app.services import payee_report as rep
from app.dates import ist_today
OWNER = "asrk.sastry"
owner = login(OWNER)
tenant = login("ten1")
sast = db.scalar(select(User).where(User.username == OWNER))
oflat = sast.flats[0]
oflat_id, oflat_no = oflat.id, oflat.number
oct_ = db.scalar(select(LedgerPeriod).where(LedgerPeriod.year == 2026, LedgerPeriod.month == 10))
def mrow(number, period=oct_):
    return db.scalar(select(MaintenancePayment).where(MaintenancePayment.period_id == period.id,
                                                      MaintenancePayment.flat_id == db.scalar(select(Flat.id).where(Flat.number == number))))
others = [n for n in db.scalars(select(Flat.number).order_by(Flat.number)).all() if n != oflat_no and mrow(n) is not None]
a_no, b_no = others[0], others[1]
a_id, b_id = mrow(a_no).id, mrow(b_no).id
owner_mp_id = mrow(oflat_no).id
for i in (a_id, b_id, owner_mp_id):
    admin.post(f"/maintenance/{i}/revert")
admin.post("/projects", data={"name": "Roof rpt", "total_cost": "25000", "due_date": (ist_today() + timedelta(days=5)).isoformat()})
db.expire_all()
proj = db.scalar(select(Project).where(Project.name == "Roof rpt"))
share = db.scalar(select(ProjectPayment).where(ProjectPayment.project_id == proj.id, ProjectPayment.flat_id == oflat_id))
today = ist_today().isoformat()

# --- under the initial UPI ID
admin.post(f"/maintenance/{a_id}/pay", data={"paid_on": "2026-10-04", "payment_note": "UTR-A-OLD"})
owner.post("/owner/maintenance", data={"flat_id": oflat_id, "period_id": oct_.id, "paid_on": "2026-10-05", "amount": "",
                                       "payment_mode": "UPI", "payment_note": "UTR123456789012"})
owner.post("/owner/projects/pay", data={"share_id": share.id, "amount": "100", "paid_on": today, "payment_mode": "Cash"})
owner.post("/owner/projects/pay", data={"share_id": share.id, "amount": "55", "paid_on": today, "payment_mode": "UPI", "payment_note": "UTR998877665544"})
db.expire_all()
rej = db.scalars(select(ProjectContribution).where(ProjectContribution.amount == 55)).first()
admin.post(f"/admin/submissions/project_contribution/{rej.id}/reject", data={"reason": "not received", "back": "/admin/actions"})
change(vpa="%s", mobile="9876543210", payee="New Payee Name", reason="Account moved", bank="50000")
# --- under the new one
admin.post(f"/maintenance/{b_id}/pay", data={"paid_on": "2026-10-06", "payment_note": "UTR-B-NEW"})
db.expire_all()
hist = db.scalars(select(PaymentDetailChange).order_by(PaymentDetailChange.id)).all()
first_id, new_id = hist[0].id, hist[-1].id
c_id = mrow(others[2]).id
admin.post(f"/maintenance/{c_id}/revert")
pend = mrow(others[3])
pend.is_paid, pend.paid_on, pend.payee_vpa, pend.payee_mobile, pend.payee_name, pend.payee_settings_id = False, None, None, None, None, None
db.commit()

def snapshot():
    db.expire_all()
    rows = {}
    for model in (MaintenancePayment, CorpusPayment, ProjectContribution, SubmissionRejection, PaymentDetailChange):
        rows[model.__tablename__] = sorted(tuple(str(getattr(r, c.name)) for c in model.__table__.columns) for r in db.scalars(select(model)).all())
    rows["audit"] = db.scalar(select(func.count()).select_from(AuditLog))
    return rows
def page_total(html):
    return (int(re.search(r'id="payee-total">₹(\d+)<', html).group(1)), int(re.search(r'id="payee-count">(\d+)<', html).group(1)))
''' % NEW


def test_filters_by_vpa_before_and_after_a_settings_change_and_totals(tmp_path):
    code = r'''
def q(**kw):
    from urllib.parse import urlencode
    return "/admin/payee-report?" + urlencode(kw)
base = admin.get("/admin/payee-report")
out["status"] = base.status_code
out["covered_note"] = ["maintenance, corpus and project" in base.text, "expenses and settlements" in base.text, "No UPI / cash-bank" in base.text]
all_rows = rep.collect_all(db)
def rows_for(**kw):
    return rep.apply_filters(all_rows, rep.parse(kw))
old = rows_for(vpa="%(OLD)s", status="all")
new = rows_for(vpa="%(NEW)s", status="all")
out["old_vpas"] = sorted({r["vpa"] for r in old})
out["new_vpas"] = sorted({r["vpa"] for r in new})
out["old_notes"] = sorted(r["ref"] for r in old if r["type"] == "maintenance")
out["new_notes"] = sorted(r["ref"] for r in new if r["type"] == "maintenance")
out["old_versions"] = sorted({r["version"] for r in old}) == [first_id]
out["new_versions"] = sorted({r["version"] for r in new}) == [new_id]
out["text_match"] = sorted({r["vpa"] for r in rows_for(vpa_text="SOCIETY", status="all")})
out["mobile"] = sorted({r["vpa"] for r in rows_for(mobile="9876543210", status="all")})
out["payee_name"] = sorted({r["vpa"] for r in rows_for(payee="new payee", status="all")})
out["by_version"] = [sorted({r["vpa"] for r in rows_for(version=str(first_id), status="all")}), sorted({r["vpa"] for r in rows_for(version=str(new_id), status="all")})]
out["date_range"] = sorted(r["ref"] for r in rows_for(date_from="2026-10-05", date_to="2026-10-06", type="maintenance"))
out["flat"] = sorted({r["flat"] for r in rows_for(flat=str(oflat_id), status="all")})
out["type_project"] = sorted({r["type"] for r in rows_for(type="project", status="all")})
out["mode_cash"] = sorted({r["mode"] for r in rows_for(mode="Cash", status="all")})

# the on-screen page applies the same filters
page = admin.get(q(vpa="%(OLD)s", status="all")).text
out["page_old_has_new"] = "UTR-B-NEW" in page
out["page_old_has_old"] = "UTR-A-OLD" in page
page = admin.get(q(vpa="%(NEW)s")).text
out["page_new"] = ["UTR-B-NEW" in page, "UTR-A-OLD" in page, page_total(page)]
out["new_total_expected"] = sum(r["amount"] for r in new if r["counted"])

# totals leave out rejected rows unless the toggle is on
full = admin.get(q(status="all")).text
total, count = page_total(full)
counted = [r for r in all_rows if r["counted"]]
out["total"] = [total, sum(r["amount"] for r in counted), count, len(counted)]
rej_rows = [r for r in all_rows if r["status"] == "rejected"]
out["rejected_rows"] = [(r["amount"], r["reason"]) for r in rej_rows]
out["rejected_marked"] = "not counted: not received" in full
inc_total, inc_count = page_total(admin.get(q(status="all", include="1")).text)
out["include_total"] = [inc_total - total, inc_count - count]
flagged = [r for r in all_rows if r["status"] in ("rejected", "reverted")]
out["rej_amount"] = sum(r["amount"] for r in flagged)
out["flagged_count"] = len(flagged)
out["statuses"] = sorted({r["status"] for r in all_rows})
only_rej = admin.get(q(status="rejected")).text
out["only_rejected"] = page_total(only_rej)

# summary per payee account and per settings period next to the recorded balance
summ = rep.build(db, rep.parse({"status": "all"}))["summary"]
by = {g["payee"]: g for g in summ["by_payee"]}
out["by_payee_keys"] = sorted(by)
out["by_payee_old_total"] = by["%(OLD)s"]["total"]
out["no_upi_total"] = by["No UPI / cash-bank"]["total"]
out["periods"] = [(p["id"], p["vpa"], p["upi_total"], p["other_total"], p["to"] is None) for p in summ["periods"]]
p0, p1 = summ["periods"][0], summ["periods"][-1]
out["period_edges"] = [p0["to"] == hist[1].changed_at, p0["fund_at_end"] == hist[1].fund_balance, p1["fund_at_start"] == hist[-1].fund_balance, p1["bank_entered"]]
out["period_text"] = "Under UPI %(NEW)s" in full and "Under UPI %(OLD)s" in full
out["period_sentence_in_ist"] = bool(re.search(r"from \d\d-\d\d-\d{4}", full)) or " IST" in full
''' % {"OLD": OLD, "NEW": NEW}
    out = run_scratch(tmp_path, SETUP + code)
    assert out["status"] == 200 and out["covered_note"] == [True, True, True]
    assert out["old_vpas"] == [OLD] and out["new_vpas"] == [NEW]
    assert "UTR-A-OLD" in out["old_notes"] and "UTR-B-NEW" not in out["old_notes"]
    assert out["new_notes"] == ["UTR-B-NEW"]
    assert out["old_versions"] is True and out["new_versions"] is True
    assert out["text_match"] == [NEW] and out["mobile"] == [NEW] and out["payee_name"] == [NEW]
    assert out["by_version"] == [[OLD], [NEW]]
    assert out["date_range"] == ["UTR-B-NEW", "UTR123456789012"]
    assert out["flat"] and len(out["flat"]) == 1
    assert out["type_project"] == ["project"] and out["mode_cash"] == ["Cash"]
    assert out["page_old_has_new"] is False and out["page_old_has_old"] is True
    assert out["page_new"][:2] == [True, False] and out["page_new"][2] == [out["new_total_expected"], 1]
    assert out["total"][0] == out["total"][1] and out["total"][2] == out["total"][3]
    assert out["rejected_rows"] == [[55, "not received"]] and out["rejected_marked"] is True
    assert out["statuses"] == ["paid", "pending", "rejected", "reverted"]
    assert out["include_total"] == [out["rej_amount"], out["flagged_count"]] and out["flagged_count"] == 2 and out["rej_amount"] == 2055
    assert out["only_rejected"] == [0, 0]
    assert OLD in out["by_payee_keys"] and NEW in out["by_payee_keys"] and "No UPI / cash-bank" in out["by_payee_keys"]
    assert out["by_payee_old_total"] > 0 and out["no_upi_total"] == 100
    assert len(out["periods"]) == 2 and out["periods"][0][1] == OLD and out["periods"][1][1] == NEW
    assert out["periods"][1][4] is True and out["periods"][0][4] is False
    assert out["periods"][1][2] == out["new_total_expected"]
    assert out["period_edges"] == [True, True, True, out["period_edges"][3]] and out["period_edges"][3]
    assert out["period_text"] is True and out["period_sentence_in_ist"] is True


def test_exports_match_the_screen_and_cover_the_filtered_view_only(tmp_path):
    code = r'''
from urllib.parse import urlencode
def check(params):
    qs = urlencode(params)
    f = rep.parse(params)
    expected = rep.apply_filters(rep.collect_all(db), f)
    page = admin.get("/admin/payee-report?" + qs)
    res = {"expected_rows": len(expected), "page_total": page_total(page.text)}
    c = admin.get("/admin/payee-report/export.csv?" + qs)
    rows = list(csv.reader(io.StringIO(c.content.decode("utf-8"))))
    res["csv"] = [c.status_code, c.headers["content-type"].split(";")[0], rows[0] == list(rep.COLUMNS), len(rows) - 1]
    cnt = rows[0].index("Counted"); amt = rows[0].index("Amount")
    res["csv_total"] = [sum(int(r[amt]) for r in rows[1:] if r[cnt] == "Yes"), sum(1 for r in rows[1:] if r[cnt] == "Yes")]
    x = admin.get("/admin/payee-report/export.xlsx?" + qs)
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    ws = wb["Payments"]
    body = [[c.value for c in row] for row in ws.iter_rows(min_row=5)]
    res["xlsx"] = [x.status_code, wb.sheetnames, len(body), [c.value for c in ws[4]] == list(rep.COLUMNS)]
    res["xlsx_total"] = sum(int(r[amt]) for r in body if r[cnt] == "Yes")
    res["xlsx_filters_line"] = ws["B2"].value
    res["xlsx_summary"] = [[c.value for c in row] for row in wb["By payee"].iter_rows()][-1]
    res["csv_times"] = rows[1][0] if rows[1:] else None
    p = admin.get("/admin/payee-report/export.pdf?" + qs)
    res["pdf"] = [p.status_code, p.content[:4].decode(), p.headers["content-type"]]
    return res
out["all"] = check({"status": "all"})
out["new"] = check({"vpa": "%(NEW)s", "status": "all"})
out["rej_included"] = check({"status": "all", "include": "1"})
out["empty"] = check({"vpa_text": "nobody@nowhere"})
out["by_flat"] = check({"flat": str(oflat_id), "status": "all"})
''' % {"NEW": NEW}
    out = run_scratch(tmp_path, SETUP + code)
    for key in ("all", "new", "rej_included", "by_flat"):
        r = out[key]
        assert r["csv"][0] == 200 and r["csv"][1] == "text/csv" and r["csv"][2] is True
        assert r["csv"][3] == r["expected_rows"] == r["xlsx"][2], key
        assert r["xlsx"][0] == 200 and r["xlsx"][1] == ["Payments", "By payee", "By settings version"] and r["xlsx"][3] is True
        assert r["csv_total"][0] == r["xlsx_total"] == r["page_total"][0], key
        assert r["csv_total"][1] == r["page_total"][1], key
        assert r["pdf"][0] == 200 and r["pdf"][1] == "%PDF" and r["pdf"][2] == "application/pdf"
    assert out["rej_included"]["page_total"][0] == out["all"]["page_total"][0] + 2055
    assert out["new"]["expected_rows"] == 1 and out["new"]["xlsx_filters_line"].count(NEW) == 1
    assert out["empty"]["expected_rows"] == 0 and out["empty"]["csv"][3] == 0 and out["empty"]["page_total"] == [0, 0]
    assert out["empty"]["pdf"][0] == 200


def test_only_admins_can_open_the_report_or_its_exports_and_the_nav_link_is_admin_only(tmp_path):
    code = r'''
urls = ["/admin/payee-report", "/admin/payee-report/export.csv", "/admin/payee-report/export.xlsx", "/admin/payee-report/export.pdf"]
out["owner"] = [owner.get(u).status_code for u in urls]
out["tenant"] = [tenant.get(u).status_code for u in urls]
anon = TestClient(app)
out["anon"] = [anon.get(u, follow_redirects=False).status_code in (302, 303, 401, 403) for u in urls]
out["admin"] = [admin.get(u).status_code for u in urls]
out["nav_admin"] = "Payee / UPI report" in admin.get("/").text
out["nav_owner"] = "Payee / UPI report" in owner.get("/").text
out["owner_leak"] = "UTR-B-NEW" in owner.get("/admin/payee-report").text
'''
    out = run_scratch(tmp_path, SETUP + code)
    assert out["owner"] == [403] * 4 and out["tenant"] == [403] * 4
    assert out["anon"] == [True] * 4 and out["admin"] == [200] * 4
    assert out["nav_admin"] is True and out["nav_owner"] is False and out["owner_leak"] is False


def test_report_is_read_only_groups_cash_as_no_upi_and_shows_empty_state(tmp_path):
    code = r'''
before = snapshot()
for u in ["/admin/payee-report", "/admin/payee-report?status=all&include=1", "/admin/payee-report?vpa=No+UPI+%2F+cash-bank&status=all",
          "/admin/payee-report?page=99", "/admin/payee-report?date_from=bad&flat=x&version=zzz&type=weird&mode=weird&status=weird",
          "/admin/payee-report/export.csv", "/admin/payee-report/export.xlsx", "/admin/payee-report/export.pdf"]:
    assert admin.get(u).status_code == 200, u
out["read_only"] = snapshot() == before
grp = rep.apply_filters(rep.collect_all(db), rep.parse({"vpa": "No UPI / cash-bank", "status": "paid"}))
out["no_upi_rows"] = sorted((r["type"], r["mode"], r["amount"], r["status"]) for r in grp)
page = admin.get("/admin/payee-report?vpa=No+UPI+%2F+cash-bank&status=paid").text
out["no_upi_page"] = ["No UPI / cash-bank" in page, "Cash" in page]
# pending dues are not listed by default but are with status=pending or everything
out["pending_default"] = any(r["status"] == "pending" for r in rep.apply_filters(rep.collect_all(db), rep.parse({})))
out["pending_listed"] = len(rep.apply_filters(rep.collect_all(db), rep.parse({"status": "pending"}))) > 0
# pagination: everything is more than one page, and the second page differs from the first
rep.PAGE_SIZE = 5
big = rep.build(db, rep.parse({"status": "all"}))
out["pages"] = [big["pages"] > 1, len(big["page_rows"]) == rep.PAGE_SIZE]
p1 = admin.get("/admin/payee-report?status=all").text
p2 = admin.get("/admin/payee-report?status=all&page=2").text
out["pager"] = ["Page 1 of" in p1, "Page 2 of" in p2, "page=2" in p1, p1 != p2]
# no matches at all: the empty state, nothing broken
empty = admin.get("/admin/payee-report?vpa_text=nobody%40nowhere")
out["empty"] = [empty.status_code, "No payments match these filters" in empty.text, "₹0" in empty.text]
# phone cards and the desktop table are both rendered
out["layouts"] = ["payee-cards" in p1, "payee-table" in p1]
# dates and times are dd-mm-yyyy / IST
out["ist"] = bool(re.search(r"\d\d-\d\d-\d{4}", p1))
'''
    out = run_scratch(tmp_path, SETUP + code)
    assert out["read_only"] is True
    assert out["no_upi_rows"] == [["project", "Cash", 100, "paid"]]
    assert out["no_upi_page"] == [True, True]
    assert out["pending_default"] is False and out["pending_listed"] is True
    assert out["pages"] == [True, True] and out["pager"] == [True, True, True, True]
    assert out["empty"] == [200, True, True]
    assert out["layouts"] == [True, True] and out["ist"] is True
